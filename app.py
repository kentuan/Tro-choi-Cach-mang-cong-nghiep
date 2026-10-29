import io
import json
import os
import random
import subprocess
import threading
import time
from datetime import datetime

from dotenv import load_dotenv
from flask import Flask, render_template, request, jsonify, send_file
from flask_socketio import SocketIO, emit, join_room
from flask_sqlalchemy import SQLAlchemy

load_dotenv()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "123456789")
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", "sqlite:///app.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)
socketio = SocketIO(app, cors_allowed_origins="*", async_mode="eventlet")
DEBUG = os.getenv("DEBUG", "False").lower() == "true"

# ── Cấu hình thời gian ─────────────────────────────────────────────

VONG1_TIME = int(os.getenv("VONG1_TIME", "90"))
VONG2_TIME = int(os.getenv("VONG2_TIME", "180"))
VONG3_TIME = int(os.getenv("VONG3_TIME", "120"))
VONG4_TIME = int(os.getenv("VONG4_TIME", "120"))
# Vòng 5 do server điều khiển từng câu; số này chỉ là giới hạn an toàn cho cả vòng
VONG5_TIME = int(os.getenv("VONG5_TIME", "600"))
# Đếm ngược trước khi mỗi vòng bắt đầu (giây) – đặt COUNTDOWN_TIME trong file .env
COUNTDOWN_TIME = max(0, int(os.getenv("COUNTDOWN_TIME", "3")))
# Thời gian chờ thêm sau khi hết giờ, đề phòng đội mất mạng không báo "đã xong"
ROUND_GRACE = 5
POINTS_PER_CORRECT = int(os.getenv("VONG1_POINTS", "5"))  # Vòng 1: điểm mỗi câu đúng (39 câu nên để thấp)
# Điểm trừ mỗi lần bấm Bỏ qua ở vòng 1–4 (chỉnh trong .env). SKIP_PENALTY áp dụng chung; muốn riêng từng vòng
# thì đặt thêm SKIP_PENALTY_VONG1 … SKIP_PENALTY_VONG4. Đặt 0 để tắt trừ điểm.
SKIP_PENALTY = int(os.getenv("SKIP_PENALTY", "5"))
SKIP_PENALTIES = {v: int(os.getenv(f"SKIP_PENALTY_VONG{v}", SKIP_PENALTY)) for v in (1, 2, 3, 4)}
POINTS_PER_PAIR = 10     # Vòng 2: điểm mỗi cặp ghép đúng
POINTS_PER_SLOT = 10     # Vòng 3: điểm mỗi vị trí xếp đúng trên dòng thời gian
POINTS_V4_RIGHT = 10     # Vòng 4: điểm mỗi thẻ phân loại đúng
POINTS_V4_WRONG = 5      # Vòng 4: điểm trừ mỗi thẻ phân loại sai (tổng điểm vòng có thể âm)

# Vòng 5 (nhanh tay & cược điểm): số câu mỗi ván, mức cược, thời lượng từng pha (giây)
V5_QUESTIONS = int(os.getenv("V5_QUESTIONS", "7"))
V5_BETS = (10, 20, 30)
V5_DEFAULT_BET = V5_BETS[0]
V5_BET_SECONDS = int(os.getenv("V5_BET_SECONDS", "5"))
V5_OPEN_SECONDS = int(os.getenv("V5_OPEN_SECONDS", "10"))
V5_REOPEN_SECONDS = 8    # mở lại chuông sau khi một đội trả lời sai
V5_ANSWER_SECONDS = int(os.getenv("V5_ANSWER_SECONDS", "8"))
V5_REVEAL_SECONDS = int(os.getenv("V5_REVEAL_SECONDS", "4"))
V6_QUESTIONS = int(os.getenv("V6_QUESTIONS", "5"))
V6_TIME = int(os.getenv("V6_TIME", "360"))
V6_BETS = (10, 20, 30, 50)
V6_RULES_SECONDS = int(os.getenv("V6_RULES_SECONDS", "30"))
V6_OPEN_SECONDS = int(os.getenv("V6_OPEN_SECONDS", "8"))
V6_DECIDE_SECONDS = int(os.getenv("V6_DECIDE_SECONDS", "8"))
V6_CALL_SECONDS = int(os.getenv("V6_CALL_SECONDS", "6"))
V6_ANSWER_SECONDS = int(os.getenv("V6_ANSWER_SECONDS", "15"))
V6_REVEAL_SECONDS = int(os.getenv("V6_REVEAL_SECONDS", "4"))

# Thời lượng từng vòng. Vòng nào có trong dict này thì BTC mới bắt đầu được.
# Thêm vòng mới: ROUND_TIME[6] = ..., rồi thêm giao diện chơi tương ứng ở client.
ROUND_TIME = {1: VONG1_TIME, 2: VONG2_TIME, 3: VONG3_TIME, 4: VONG4_TIME, 5: VONG5_TIME, 6: V6_TIME}


# ── Database Models ─────────────────────────────────────────────────
class Score(db.Model):
    __tablename__ = "scores"
    id = db.Column(db.Integer, primary_key=True)
    group = db.Column(db.String(30), nullable=False)
    team_name = db.Column(db.String(30), nullable=False)
    vong = db.Column(db.Integer, nullable=False, default=1)
    diem = db.Column(db.Integer, nullable=False, default=0)

    def to_dict(self):
        return {
            "id": self.id,
            "group": self.group,
            "team_name": self.team_name,
            "vong": self.vong,
            "diem": self.diem,
        }


# ── Trạng thái trò chơi (lưu trong RAM) ────────────────────────────
ALLOWED_GROUPS = ["Nhóm 1", "Nhóm 2", "Nhóm 4", "Nhóm 5"]

# Thông tin đội đã tham gia:
# { "Nhóm 1": {"team_name": "...", "sid": "...", "confirmed": False}, ... }
teams = {}
btc_sids = set()

# status:
#   "waiting"   : chưa có vòng nào chạy (trước vòng 1 hoặc giữa các vòng)
#   "countdown" : đang đếm ngược trước khi vào vòng
#   "playing"   : vòng đang diễn ra
# round : số của vòng gần nhất đã bắt đầu (0 = chưa có vòng nào)
# run_id: tăng mỗi khi vòng kết thúc / reset để huỷ các bộ đếm giờ cũ
game_state = {"status": "waiting", "round": 0, "run_id": 0, "round_deadline": 0.0, "countdown_deadline": 0.0}

# BTC manually selects what the projector displays.
tv_control = {"view": "welcome", "round": 1}

# Điểm trên TV chỉ hiện sau khi BTC cho phép, theo từng đội:
# { vong (int) hoặc "total": {"Nhóm 1", ...} } – những đội đã được phép hiện điểm trên TV.
# Tách riêng với revealed_rounds (công bố điểm cho thiết bị của các đội chơi).
tv_revealed = {}

# Các nhóm đã báo hoàn thành vòng hiện tại
finished_groups = set()

# Các vòng mà BTC đã công bố điểm cho đội chơi (mỗi vòng công bố riêng)
revealed_rounds = set()

# Vòng 2, 3: các cặp/vị trí đã được tính điểm, dạng (vòng, nhóm, số thứ tự câu, số thứ tự cặp/vị trí)
# – để mỗi mục chỉ tính một lần dù client gửi lặp
awarded_items = set()

# Vòng 1–4: các lần Bỏ qua đã bị trừ điểm, dạng (vòng, nhóm, số thứ tự câu/bảng/thẻ) – mỗi mục chỉ trừ một lần
skipped_items = set()

# Vòng 4: kết quả lần trả lời đầu tiên của mỗi thẻ, dạng {(nhóm, số thứ tự thẻ): đúng/sai}
v4_results = {}

# Vòng 5: trạng thái câu hỏi đang chạy. Các handler socket và tác vụ nền cùng đọc/ghi nên mọi thay đổi
# đi qua v5_lock (chỉ giữ khóa cho thao tác dict ngắn; không gửi tin, không ghi DB, không sleep khi đang giữ).
v5_lock = threading.Lock()
v5 = {
    "questions": [],   # các câu của ván này: {chu_de, cau_hoi, options, dung, trang, giai_thich}
    "q": -1,           # chỉ số câu hiện tại
    "phase": "idle",   # idle | prep | open | answering | reveal
    "span": 1,         # độ dài pha hiện tại (giây), để client vẽ vòng đếm giờ
    "deadline": 0.0,   # hạn chót pha hiện tại (time.monotonic)
    "bets": {},        # {nhóm: mức cược của câu hiện tại}
    "totals": {},      # {nhóm: tổng điểm vòng 5}
    "locked": set(),   # các đội đã trả lời ở câu hiện tại (hết quyền)
    "buzzer": None,    # đội đang giữ chuông
    "attempts": [],    # các lượt trả lời của câu hiện tại
    "winner": None,    # đội trả lời đúng ở câu hiện tại
}

# Vòng 6 chạy hai trận cùng nhịp; mọi thao tác và điểm đều được xác nhận ở server.
v6_lock = threading.Lock()
v6 = {"matches": [], "questions": [], "q": -1, "phase": "idle", "span": 1,
      "deadline": 0.0, "final_revealed": False}


# ── Helper ───────────────────────────────────────────────────────────
def _round_stats(vong, group):
    """(số đúng, số sai) của một đội ở vòng đang chơi – tính từ dữ liệu server, không tin số client báo."""
    if vong == 4:
        mine = [ok for (g, _), ok in v4_results.items() if g == group]
        return sum(mine), len(mine) - sum(mine)
    return sum(1 for k in awarded_items if k[0] == vong and k[1] == group), 0


def _load_json(name):
    """Đọc file dữ liệu trong static/games (đọc lại mỗi lần để sửa đề không cần khởi động lại)."""
    with open(os.path.join(app.static_folder, "games", name), "r", encoding="utf-8") as f:
        return json.load(f)


def _game_locked():
    """Đã bắt đầu chơi (kể cả giữa các vòng): không cho vào/ra nhóm nữa."""
    return game_state["status"] != "waiting" or game_state["round"] > 0


def _next_round():
    """Vòng tiếp theo BTC có thể bắt đầu, hoặc None."""
    if game_state["status"] != "waiting":
        return None
    nxt = game_state["round"] + 1
    return nxt if nxt in ROUND_TIME else None


def _scoreboard(vong):
    """Bảng điểm một vòng: mọi đội đang tham gia, điểm cao xếp trước."""
    rows = {s.group: s.diem for s in Score.query.filter_by(vong=vong).all()}
    board = [
        {"group": g, "team_name": teams[g]["team_name"], "diem": rows.get(g, 0)}
        for g in ALLOWED_GROUPS
        if g in teams
    ]
    board.sort(key=lambda r: -r["diem"])
    return board


TV_CONTENT = {
    "rules": {
        1: ("Luật chơi · Vòng 1", "Sắp xếp các chữ cái thành đáp án đúng. Có thể bấm thẻ để điền ô trống tiếp theo hoặc kéo thẻ vào ô. Bấm Kiểm tra để chốt; Bỏ qua bị trừ điểm theo cấu hình."),
        2: ("Luật chơi · Vòng 2", "Ghép từng thẻ đáp án với gợi ý tương ứng. Có thể bấm thẻ để điền ô trống kế tiếp hoặc kéo thẻ. Hoàn thành các cặp rồi bấm Kiểm tra."),
        3: ("Luật chơi · Vòng 3", "Xếp các thẻ theo đúng thứ tự thời gian. Bấm thẻ để điền vị trí trống tiếp theo hoặc kéo thẻ vào vị trí muốn đặt, sau đó bấm Kiểm tra."),
        4: ("Luật chơi · Vòng 4", "Đọc thẻ sự kiện và chọn ô phân loại phù hợp. Mỗi thẻ chỉ được trả lời một lần; trả lời đúng được cộng điểm, sai bị trừ điểm."),
        5: ("Luật chơi · Vòng 5", "Chọn mức cược khi chỉ thấy chủ đề. Khi câu hỏi mở, bấm chuông để giành quyền trả lời. Đúng được cộng cược, sai bị trừ cược; điểm có thể âm."),
        6: ("Luật chơi · Vòng 6", "Hai trận đối đầu diễn ra song song. Đội mở cược, đội kia chọn Theo, Tố hoặc Bỏ. Hai đội trả lời riêng; câu cuối nhân đôi mức cược. Nếu hoà, có câu phụ."),
    },
    "guide": {
        1: ("Hướng dẫn · Vòng 1", "Chạm hoặc bấm một chữ cái để đưa thẻ vào ô trống kế tiếp. Kéo thẻ nếu muốn đặt ở vị trí cụ thể. Dùng Bỏ qua nếu cần chuyển câu."),
        2: ("Hướng dẫn · Vòng 2", "Bấm thẻ trong khay để ghép vào gợi ý trống tiếp theo. Bấm thẻ đã đặt để trả về khay; kéo thẻ để đổi vị trí."),
        3: ("Hướng dẫn · Vòng 3", "Bấm thẻ để đưa vào vị trí trống tiếp theo, hoặc kéo để sắp xếp. Khi đã xếp đủ, kiểm tra thứ tự."),
        4: ("Hướng dẫn · Vòng 4", "Bấm một trong các ô phân loại để xếp thẻ đang hiện. Đọc giải thích sau mỗi lượt để biết kết quả."),
        5: ("Hướng dẫn · Vòng 5", "Trong pha cược, chọn mức phù hợp. Khi câu hỏi hiện ra, bấm BẤM CHUÔNG; nếu giành quyền, chọn đáp án trước khi hết giờ."),
        6: ("Hướng dẫn · Vòng 6", "Theo dõi vai trò của đội mình trên thiết bị. Chọn cược hoặc Theo/Tố/Bỏ đúng pha; ở pha trả lời, hai đội chọn độc lập."),
    },
}


def _tv_payload():
    view, vong = tv_control["view"], tv_control["round"]
    if view in TV_CONTENT:
        title, body = TV_CONTENT[view].get(vong, ("", ""))
    elif view == "time":
        title, body = f"Thời gian · Vòng {vong}", ""
    elif view == "scores":
        title, body = ("Tổng điểm", "") if vong == "total" else (f"Điểm · Vòng {vong}", "")
    else:
        title, body = "Trò chơi Cách mạng Công nghiệp", "Chào mừng các đội chơi!"

    active = game_state["status"] == "playing" and game_state["round"] == vong
    countdown = game_state["status"] == "countdown" and game_state["round"] == vong
    if active:
        seconds = max(0, game_state["round_deadline"] - time.monotonic())
    elif countdown:
        seconds = max(0, game_state["countdown_deadline"] - time.monotonic())
    elif game_state["status"] == "waiting" and game_state["round"] == vong and game_state["round"] > 0:
        seconds = 0
    else:
        seconds = ROUND_TIME.get(vong, 0)
    scoreboard, hidden_count = [], 0
    if view == "scores":
        full = _final_board() if vong == "total" else _scoreboard(vong) if isinstance(vong, int) else []
        shown = tv_revealed.get(vong, set())
        # Chỉ gồm các đội BTC đã cho hiện; thứ hạng tính trong số đội này
        scoreboard = [r for r in full if r["group"] in shown]
        hidden_count = len(full) - len(scoreboard)
    payload = {
        "view": view, "round": vong, "title": title, "body": body,
        "game_status": game_state["status"], "active_round": game_state["round"],
        "active": active, "countdown": countdown, "seconds": seconds,
        "scoreboard": scoreboard,
        "hidden_count": hidden_count,
        "final_revealed": bool(v6.get("final_revealed")),
    }
    return payload


def _tv_control_payload():
    return {
        "view": tv_control["view"], "round": tv_control["round"],
        # {"1": ["Nhóm 1", ...], "total": [...]}: đội nào đang được hiện điểm trên TV
        "revealed": {str(k): sorted(v) for k, v in tv_revealed.items()},
    }


def _broadcast_tv():
    socketio.emit("tv_update", _tv_payload(), to="tv")
    socketio.emit("tv_control", _tv_control_payload(), to="btc")


def _state_payload():
    return {
        "teams": {g: t["team_name"] for g, t in teams.items()},
        "taken_groups": list(teams.keys()),
        "confirmed": {g: t["confirmed"] for g, t in teams.items()},
        "game_status": game_state["status"],
        "round": game_state["round"],
        "next_round": _next_round(),
        "finished": sorted(finished_groups),
        # Chỉ gửi điểm của những vòng BTC đã công bố
        "revealed": {str(v): _scoreboard(v) for v in sorted(revealed_rounds)},
        "final_revealed": bool(v6.get("final_revealed")),
        "final_scores": _final_board() if v6.get("final_revealed") else [],
    }


def _broadcast_state():
    """Phát trạng thái cho tất cả client (đội chơi + BTC)."""
    socketio.emit("state_update", _state_payload())
    _broadcast_tv()


def _btc_scores_payload():
    return {
        "rounds": [
            {"vong": v, "revealed": v in revealed_rounds, "scores": _scoreboard(v)}
            for v in range(1, game_state["round"] + 1)
        ],
        "final_revealed": bool(v6.get("final_revealed")),
        "final": _final_board() if v6.get("final_revealed") else [],
        # BTC luôn thấy tổng điểm trực tiếp (để chọn đội hiện trên TV), kể cả khi chưa công bố cho đội chơi
        "live_final": _final_board() if game_state["round"] >= 6 else [],
        "v6_pairings": _v6_pairings() if game_state["round"] >= 5 else [],
    }


def _final_board():
    rows = []
    for g, team in teams.items():
        rounds = {}
        for v in range(1, 7):
            score = Score.query.filter_by(group=g, vong=v).first()
            rounds[v] = score.diem if score else 0
        rows.append({"group": g, "team_name": team["team_name"], "diem": sum(rounds.values()), "rounds": rounds})
    rows.sort(key=lambda r: (-r["diem"], -r["rounds"][6], -r["rounds"][5], -r["rounds"][4]))
    return rows


def _emit_btc_scores():
    """Điểm trực tiếp chỉ gửi cho BTC (phòng 'btc'), đội chơi không nhận được."""
    socketio.emit("btc_scores", _btc_scores_payload(), to="btc")
    if tv_control["view"] == "scores":
        _broadcast_tv()


def _pending_groups():
    """Các nhóm đã vào nhưng chưa xác nhận tên (theo thứ tự nhóm)."""
    return [g for g in ALLOWED_GROUPS if g in teams and not teams[g]["confirmed"]]


def _end_round():
    """Kết thúc vòng hiện tại. Gọi trong app context."""
    if game_state["status"] != "playing":
        return
    game_state["status"] = "waiting"
    game_state["round_deadline"] = 0.0
    game_state["run_id"] += 1  # huỷ bộ đếm giờ đang chờ
    socketio.emit("round_ended", {"round": game_state["round"]})
    _broadcast_state()
    _emit_btc_scores()


def _run_round(run_id, vong):
    """Chạy nền: đếm ngược → bắt đầu vòng → tự kết thúc khi hết giờ."""
    socketio.sleep(COUNTDOWN_TIME)
    if game_state["run_id"] != run_id:  # đã bị reset trong lúc đếm ngược
        return

    with app.app_context():
        game_state["status"] = "playing"
        game_state["countdown_deadline"] = 0.0
        game_state["round_deadline"] = time.monotonic() + ROUND_TIME[vong]
        socketio.emit("game_started", {"round": vong, "time": ROUND_TIME[vong]})
        _broadcast_state()

    if vong == 5:
        # Vòng 5: server chạy từng câu; ROUND_TIME[5] chỉ là giới hạn an toàn cho cả vòng
        _run_quiz(run_id, time.monotonic() + ROUND_TIME[vong])
    elif vong == 6:
        with app.app_context():
            _run_v6(run_id, time.monotonic() + ROUND_TIME[vong])
    else:
        socketio.sleep(ROUND_TIME[vong] + ROUND_GRACE)
    if game_state["run_id"] != run_id:  # vòng đã kết thúc sớm hoặc bị reset
        return

    with app.app_context():
        _end_round()


# ── Vòng 5: nhanh tay & cược điểm (server điều khiển từng câu) ──────
def _v5_reset():
    with v5_lock:
        v5.update(questions=[], q=-1, phase="idle", span=1, deadline=0.0, bets={}, totals={},
                  locked=set(), buzzer=None, attempts=[], winner=None)


def _v5_team_name(group):
    return teams.get(group, {}).get("team_name", group)


def _v5_set_phase(phase, span):
    """Đổi pha và đặt hạn chót. Gọi khi ĐANG giữ v5_lock."""
    v5["phase"] = phase
    v5["span"] = span
    v5["deadline"] = time.monotonic() + span


def _v5_load_questions():
    """Chọn ngẫu nhiên V5_QUESTIONS câu từ vong_5.json; xáo thứ tự đáp án (trừ câu Đúng/Sai)."""
    pool = _load_json("vong_5.json")
    questions = []
    for item in random.sample(pool, min(V5_QUESTIONS, len(pool))):
        order = list(range(len(item["lua_chon"])))
        if len(order) > 2:
            random.shuffle(order)
        questions.append({
            "chu_de": item["chu_de"],
            "cau_hoi": item["cau_hoi"],
            "options": [item["lua_chon"][i] for i in order],
            "dung": order.index(item["dung"]),
            "trang": item.get("trang"),
            "giai_thich": item.get("giai_thich", ""),
        })
    return questions


def _v5_snapshot(group=None):
    """Trạng thái vòng 5 để gửi cho client. Truyền group để thêm phần riêng của đội đó:
    mức cược, điểm, và các lựa chọn nếu đội đó đang giữ chuông. Nội dung câu hỏi chỉ có từ pha 'open'."""
    with v5_lock:
        q = v5["q"]
        total = len(v5["questions"])
        snap = {
            "phase": v5["phase"],
            "q": q + 1 if q >= 0 else 0,
            "total": total,
            "seconds": max(0.0, v5["deadline"] - time.monotonic()),
            "span": v5["span"],
            "locked": sorted(v5["locked"]),
            "attempts": [dict(x) for x in v5["attempts"]],
        }
        cur = v5["questions"][q] if 0 <= q < total else None
        buzzer, winner = v5["buzzer"], v5["winner"]
        bet = v5["bets"].get(group, V5_DEFAULT_BET) if group else None
        diem = v5["totals"].get(group, 0) if group else None

    if cur:
        snap["chu_de"] = cur["chu_de"]
        if snap["phase"] == "prep":
            snap["bets"] = list(V5_BETS)
        else:
            snap["cau_hoi"] = cur["cau_hoi"]
        if buzzer:
            snap["buzzer"] = {"group": buzzer, "team_name": _v5_team_name(buzzer)}
        if snap["phase"] == "reveal":
            snap["dap_an"] = cur["options"][cur["dung"]]
            snap["trang"] = cur["trang"]
            snap["giai_thich"] = cur["giai_thich"]
            if winner:
                snap["winner"] = {"group": winner, "team_name": _v5_team_name(winner)}
        if group and snap["phase"] == "answering" and buzzer == group:
            snap["options"] = list(cur["options"])
    if group:
        snap["bet"] = bet
        snap["diem"] = diem
    return snap


def _v5_broadcast():
    socketio.emit("v5_state", _v5_snapshot())


def _v5_add_score(group, delta):
    """Cộng/trừ thẳng vào điểm vòng 5 hiện có (có thể âm). Trả về tổng mới."""
    with app.app_context():
        score = Score.query.filter_by(group=group, vong=5).first()
        if score:
            score.diem += delta
            score.team_name = _v5_team_name(group)
        else:
            score = Score(group=group, team_name=_v5_team_name(group), vong=5, diem=delta)
            db.session.add(score)
        db.session.commit()
        total = score.diem
        _emit_btc_scores()
    with v5_lock:
        v5["totals"][group] = total
    return total


def _v5_to_reveal(winner):
    """Chuyển sang hiện đáp án. Trả True nếu vừa chuyển (False nếu đã ở pha này rồi)."""
    with v5_lock:
        if v5["phase"] == "reveal":
            return False
        v5["winner"] = winner
        _v5_set_phase("reveal", V5_REVEAL_SECONDS)
    return True


def _v5_resolve_answer(group, choice):
    """Chấm câu trả lời của đội đang giữ chuông (choice=None: hết giờ, tính là sai).
    Đúng: kết thúc câu. Sai: đội đó mất quyền, mở lại chuông cho các đội còn lại (hết đội thì hiện đáp án).
    Trả về False nếu không còn hợp lệ (đã có luồng khác xử lý trước)."""
    with v5_lock:
        if v5["phase"] != "answering" or v5["buzzer"] != group:
            return False
        cur = v5["questions"][v5["q"]]
        correct = choice is not None and choice == cur["dung"]
        bet = v5["bets"].get(group, V5_DEFAULT_BET)
        delta = bet if correct else -bet
        qn, total = v5["q"] + 1, len(v5["questions"])
        v5["attempts"].append({"group": group, "team_name": _v5_team_name(group), "correct": correct})
        v5["locked"].add(group)
        v5["buzzer"] = None
        if correct:
            v5["winner"] = group
            _v5_set_phase("reveal", V5_REVEAL_SECONDS)
        elif v5["locked"] >= set(teams):
            _v5_set_phase("reveal", V5_REVEAL_SECONDS)
        else:
            _v5_set_phase("open", V5_REOPEN_SECONDS)

    diem = _v5_add_score(group, delta)
    _v5_broadcast()
    # Kết quả riêng gửi vào phòng của đội (các đội khác không thấy điểm); BTC thấy nhật ký
    socketio.emit("v5_result", {"q": qn, "correct": correct, "delta": delta, "diem": diem,
                                "timeout": choice is None}, to="grp:" + group)
    socketio.emit("btc_v5", {"q": qn, "total": total, "group": group, "team_name": _v5_team_name(group),
                             "correct": correct, "delta": delta, "bet": bet, "timeout": choice is None}, to="btc")
    return True


def _v5_alive(run_id, hard_deadline):
    return (game_state["run_id"] == run_id and game_state["status"] == "playing"
            and time.monotonic() < hard_deadline)


def _v5_sleep_until(run_id, hard_deadline, until):
    """Ngủ (nhịp 0.1s) tới mốc `until`. Trả False nếu vòng đã bị huỷ/reset hoặc quá giới hạn thời gian."""
    while time.monotonic() < until:
        if not _v5_alive(run_id, hard_deadline):
            return False
        socketio.sleep(0.1)
    return _v5_alive(run_id, hard_deadline)


def _run_quiz(run_id, hard_deadline):
    """Chạy nền: lần lượt từng câu (cược → bấm chuông/trả lời → hiện đáp án). Hết câu thì trả về,
    _run_round sẽ kết thúc vòng bằng _end_round như các vòng khác."""
    questions = _v5_load_questions()
    with v5_lock:
        v5.update(questions=questions, q=-1, phase="idle", bets={}, locked=set(),
                  buzzer=None, attempts=[], winner=None)

    for qi in range(len(questions)):
        # 1) cược
        with v5_lock:
            v5.update(q=qi, bets={}, locked=set(), buzzer=None, attempts=[], winner=None)
            _v5_set_phase("prep", V5_BET_SECONDS)
            until = v5["deadline"]
        _v5_broadcast()
        if not _v5_sleep_until(run_id, hard_deadline, until):
            return


        # 2) mở chuông
        with v5_lock:
            _v5_set_phase("open", V5_OPEN_SECONDS)
        _v5_broadcast()

        # 3) chờ chuông / câu trả lời / hết giờ cho tới khi sang pha hiện đáp án
        while True:
            if not _v5_sleep_until(run_id, hard_deadline, time.monotonic() + 0.1):
                return
            with v5_lock:
                phase, deadline, buzzer = v5["phase"], v5["deadline"], v5["buzzer"]
            if phase == "reveal":
                break
            if time.monotonic() >= deadline:
                if phase == "open":
                    if _v5_to_reveal(None):  # không đội nào bấm / mọi đội còn lại không bấm kịp
                        _v5_broadcast()
                elif phase == "answering":
                    _v5_resolve_answer(buzzer, None)  # đội giữ chuông hết giờ = sai

        # 4) hiện đáp án
        with v5_lock:
            until = v5["deadline"]
        if not _v5_sleep_until(run_id, hard_deadline, until):
            return


# ── Vòng 6: poker kiến thức, hai trận song song ─────────────────────
def _v6_total(group):
    return sum(s.diem for s in Score.query.filter_by(group=group).all())


def _v6_order():
    def key(g):
        pts = {v: (Score.query.filter_by(group=g, vong=v).first() or Score(diem=0)).diem for v in range(1, 6)}
        return (-sum(pts.values()), -pts[5], -pts[4], -pts[3])
    return sorted(teams, key=key)


def _v6_load_questions():
    pool = _load_json("vong_6.json")
    normal = [item for item in pool if item.get("kind") not in ("final", "tiebreak")]
    final = [item for item in pool if item.get("kind") == "final"]
    tie = [item for item in pool if item.get("kind") == "tiebreak"]
    picked = random.sample(normal, min(V6_QUESTIONS, len(normal))) + final[:1] + tie[:1]
    out = []
    for item in picked:
        correct = item["lua_chon"][item["dung"]]
        # Đáp án sai gốc + đáp án sai bổ sung (không bắt buộc): mỗi đội sẽ được rút bộ riêng
        wrong_pool = list(dict.fromkeys(w for w in item["lua_chon"] + item.get("lua_chon_sai", []) if w != correct))
        out.append({"chu_de": item["chu_de"], "cau_hoi": item["cau_hoi"],
                    "correct": correct, "wrong_pool": wrong_pool, "n_options": len(item["lua_chon"]),
                    "options": list(item["lua_chon"]), "dung": item["dung"],
                    "trang": item.get("trang"), "giai_thich": item.get("giai_thich", ""),
                    "kind": item.get("kind", "normal")})
    return out


def _v6_variant(question, avoid=None):
    """Bộ đáp án riêng của một đội: cùng câu hỏi và đáp án đúng, nhưng thứ tự và các đáp án sai khác nhau."""
    wrong_count = min(question["n_options"] - 1, len(question["wrong_pool"]))
    for _ in range(10):
        options = [question["correct"]] + random.sample(question["wrong_pool"], wrong_count)
        random.shuffle(options)
        if options != avoid:
            break
    return {"options": options, "dung": options.index(question["correct"])}


def _v6_make_variants(match, question):
    """Mỗi đội trong trận một bộ đáp án; cố gắng không trùng với đội kia."""
    a, b = match["groups"]
    va = _v6_variant(question)
    return {a: va, b: _v6_variant(question, avoid=va["options"])}


def _v6_correct(match, group, question):
    variant = match.get("variants", {}).get(group)
    answer = match["answers"].get(group)
    return answer is not None and variant is not None and answer == variant["dung"]


def _v6_pairings():
    order = _v6_order()
    return [[order[i + 1], order[i]] for i in range(0, len(order) - 1, 2)]


def _v6_snapshot(group=None):
    with v6_lock:
        q, phase = v6["q"], v6["phase"]
        result = {"q": q + 1, "total": max(V6_QUESTIONS + 1, q + 1), "phase": phase,
                  "seconds": max(0, v6["deadline"] - time.monotonic()), "span": v6["span"], "matches": []}
        question = v6["questions"][q] if 0 <= q < len(v6["questions"]) else None
        for m in v6["matches"]:
            row = {"groups": m["groups"], "names": [_v5_team_name(g) for g in m["groups"]],
                   "scores": dict(m["scores"]), "opener": m["opener"], "bet": m["bet"],
                   "choice": m["choice"], "call": m.get("call"), "answered": len(m["answers"]), "winner": m.get("winner"),
                   "folded": m.get("folded"), "tie": m.get("tie", False), "tiebreak_skip": m.get("tiebreak_skip", False)}
            if question:
                row["chu_de"] = question["chu_de"]
                if phase in ("answer", "reveal"):
                    row["cau_hoi"] = question["cau_hoi"]
                if phase == "reveal":
                    row.update(dap_an=question["correct"], trang=question["trang"], giai_thich=question["giai_thich"],
                               answers=dict(m["answers"]),
                               # kết quả từng đội tính theo bộ đáp án riêng của đội đó
                               results={g: (None if m["answers"].get(g) is None else _v6_correct(m, g, question)) for g in m["groups"]})
            result["matches"].append(row)
        if group:
            result["mine"] = next((i for i, m in enumerate(v6["matches"]) if group in m["groups"]), None)
            result["group"] = group
            result["bets"] = [20, 40, 60, 100] if q == V6_QUESTIONS else list(V6_BETS)
            if result["mine"] is not None and phase == "answer":
                mine = v6["matches"][result["mine"]]
                answers = mine["answers"]
                result["options"] = list(mine["variants"][group]["options"])
                result["my_answered"] = group in answers
                result["my_choice"] = answers.get(group)  # để client đánh dấu đáp án đã chọn
        return result


def _v6_broadcast():
    for m in v6["matches"]:
        for g in m["groups"]:
            socketio.emit("v6_state", _v6_snapshot(g), to="grp:" + g)
    socketio.emit("v6_state", _v6_snapshot(), to="btc")


def _v6_set_phase(phase, seconds):
    with v6_lock:
        v6["phase"] = phase
        v6["span"] = seconds
        v6["deadline"] = time.monotonic() + seconds


def _v6_award(group, delta):
    score = Score.query.filter_by(group=group, vong=6).first()
    if score is None:
        score = Score(group=group, team_name=_v5_team_name(group), vong=6, diem=0)
        db.session.add(score)
    score.diem += delta
    db.session.commit()
    _emit_btc_scores()


def _v6_resolve():
    question = v6["questions"][v6["q"]]
    for m in v6["matches"]:
        if m.get("tiebreak_skip"):
            continue
        a, b = m["groups"]
        if m.get("choice") == "fold" or m.get("call") == "fold":
            winner = m["opener"] if m.get("call") == "fold" else m["opener"]
            if m.get("choice") == "fold": winner = m["opener"]
            else: winner = b if m["opener"] == a else a
            loser = b if winner == a else a
            m["scores"][winner] += 10; m["scores"][loser] -= 10
            m["winner"] = winner
            continue
        correct = {g: _v6_correct(m, g, question) for g in (a, b)}
        stake = m["bet"]
        if m.get("call") == "fold":
            loser = m["opener"]
            winner = b if loser == a else a
            m["scores"][loser] -= 10; m["scores"][winner] += 10
            m["winner"] = winner
            continue
        if correct[a] and not correct[b]: m["scores"][a] += stake; m["scores"][b] -= stake
        elif correct[b] and not correct[a]: m["scores"][b] += stake; m["scores"][a] -= stake
        elif correct[a] and correct[b]:
            speed = m.get("speed", {})
            if speed:
                fastest = min(speed, key=speed.get); m["scores"][fastest] += 5
        m["winner"] = a if correct[a] and not correct[b] else b if correct[b] and not correct[a] else None
    for m in v6["matches"]:
        for g, delta in m["scores"].items():
            # scores are cumulative; persist only this question's difference via per-question ledger
            prev = m.setdefault("persisted", {}).get(g, 0)
            _v6_award(g, delta - prev)
            m["persisted"][g] = delta


def _run_v6(run_id, hard_deadline):
    questions = _v6_load_questions()
    with v6_lock:
        v6.update(questions=questions, q=-1, phase="idle", final_revealed=False, span=1, deadline=0)
        v6["matches"] = [{"groups": pair, "scores": {g: 0 for g in pair}, "opener": pair[0], "bet": 10,
                           "choice": None, "answers": {}, "speed": {}, "persisted": {g: 0 for g in pair}}
                          for pair in _v6_pairings()]
    _v6_set_phase("rules", V6_RULES_SECONDS); _v6_broadcast()
    socketio.sleep(V6_RULES_SECONDS)
    for qi in range(min(V6_QUESTIONS + 1, len(questions))):
        with v6_lock:
            v6["q"] = qi
            for m in v6["matches"]:
                m.update(choice=None, answers={}, speed={}, call=None, folded=None, winner=None, tie=False, base_bet=None,
                         variants=_v6_make_variants(m, v6["questions"][qi]))
                if qi == V6_QUESTIONS:
                    a, b = m["groups"]
                    m["opener"] = min((a, b), key=lambda g: (m["scores"][g], -_v6_order().index(g)))
                    m["bet"] = 20
                else:
                    m["opener"] = m["groups"][qi % 2] if qi else m["groups"][0]
                    m["bet"] = 10
        _v6_set_phase("bet", V6_OPEN_SECONDS); _v6_broadcast()
        socketio.sleep(V6_OPEN_SECONDS)
        _v6_set_phase("decide", V6_DECIDE_SECONDS); _v6_broadcast()
        socketio.sleep(V6_DECIDE_SECONDS)
        with v6_lock:
            needs_call = any(m.get("choice") == "raise" for m in v6["matches"])
        if needs_call:
            _v6_set_phase("call", V6_CALL_SECONDS); _v6_broadcast()
            socketio.sleep(V6_CALL_SECONDS)
        with v6_lock:
            needs_answer = any(m.get("choice") != "fold" and m.get("call") != "fold" for m in v6["matches"])
        if needs_answer:
            _v6_set_phase("answer", V6_ANSWER_SECONDS); _v6_broadcast()
            socketio.sleep(V6_ANSWER_SECONDS)
        _v6_resolve()
        _v6_set_phase("reveal", V6_REVEAL_SECONDS); _v6_broadcast()
        socketio.sleep(V6_REVEAL_SECONDS)
        if game_state["run_id"] != run_id or time.monotonic() >= hard_deadline: return
    ties = [m for m in v6["matches"] if m["scores"][m["groups"][0]] == m["scores"][m["groups"][1]]]
    if ties and len(questions) > V6_QUESTIONS + 1:
        with v6_lock:
            v6["q"] = V6_QUESTIONS + 1
            for m in v6["matches"]:
                m.update(choice="follow", call="follow", bet=20, answers={}, speed={}, tiebreak_skip=m not in ties)
        _v6_set_phase("answer", V6_ANSWER_SECONDS); _v6_broadcast()
        socketio.sleep(V6_ANSWER_SECONDS)
        _v6_resolve()
        _v6_set_phase("reveal", V6_REVEAL_SECONDS); _v6_broadcast()
        socketio.sleep(V6_REVEAL_SECONDS)
    _end_round()

# ── Trang cho đội chơi ──────────────────────────────────────────────
@app.route("/")
def index():
    # Toàn bộ luồng chơi (chọn nhóm → đặt tên → chờ → Vòng 1) nằm trong một trang
    return render_template("index.html", vong1_time=VONG1_TIME, vong2_time=VONG2_TIME, vong3_time=VONG3_TIME, vong4_time=VONG4_TIME)


# ── Trang cho BTC ───────────────────────────────────────────────────
@app.route("/btc")
def btc_dashboard():
    return render_template("btc.html", debug=DEBUG, rounds=sorted(ROUND_TIME))


# ── Xuất kết quả ra Excel ───────────────────────────────────────────
EXPORT_ROUNDS = range(1, 7)  # vòng 1..6


def summarize(raw, teams=None):
    """Gộp các dòng điểm thô thành bảng xếp hạng.

    raw   : [(id, group, team_name, vong, diem), ...]
    teams : {group: team_name} các đội đang tham gia; đội chưa có điểm vẫn có dòng (0 điểm).
    """
    board = {}

    def team_row(group, name):
        return board.setdefault(group, {"group": group, "team_name": name, "rounds": {v: 0 for v in EXPORT_ROUNDS}, "seen": set()})

    for _id, group, team_name, vong, diem in raw:
        team = team_row(group, team_name)
        team["team_name"] = team_name or team["team_name"]
        # App dùng .first() cho mỗi (đội, vòng) nên chỉ lấy dòng đầu tiên của mỗi cặp
        if vong in team["rounds"] and vong not in team["seen"]:
            team["rounds"][vong] = diem
            team["seen"].add(vong)
    for group, name in (teams or {}).items():
        team_row(group, name)["team_name"] = name  # tên đội hiện tại trong app là tên chuẩn

    rows = list(board.values())
    for team in rows:
        team["total"] = sum(team["rounds"].values())
    # Cùng luật xếp hạng như bảng tổng kết trong app: tổng, rồi vòng 6, 5, 4
    rows.sort(key=lambda r: (-r["total"], -r["rounds"][6], -r["rounds"][5], -r["rounds"][4], r["group"]))
    return rows


def build_workbook(raw, board):
    """Tạo Workbook (openpyxl) gồm 2 sheet: Tổng kết và Chi tiết."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border

    def autosize(ws, minimum=8, maximum=40):
        for col in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 3
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(minimum, min(maximum, width))

    # Sheet 1: tổng kết
    ws = wb.active
    ws.title = "Tổng kết"
    ws.append(["Hạng", "Nhóm", "Tên đội"] + [f"Vòng {v}" for v in EXPORT_ROUNDS] + ["Tổng"])
    style_header(ws)
    for rank, team in enumerate(board, start=1):
        ws.append([rank, team["group"], team["team_name"]] + [team["rounds"][v] for v in EXPORT_ROUNDS] + [team["total"]])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            if cell.column != 3:
                cell.alignment = center
        row[-1].font = Font(bold=True)
    ws.freeze_panes = "D2"
    autosize(ws)

    # Sheet 2: chi tiết từng dòng điểm trong database
    detail = wb.create_sheet("Chi tiết")
    detail.append(["ID", "Nhóm", "Tên đội", "Vòng", "Điểm"])
    style_header(detail)
    for row in raw:
        detail.append(list(row))
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            if cell.column != 3:
                cell.alignment = center
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    autosize(detail)
    return wb


@app.route("/btc/export.xlsx")
def btc_export_excel():
    """BTC tải kết quả (điểm từng vòng + xếp hạng + chi tiết) dưới dạng Excel. Dùng được bất kỳ lúc nào."""
    raw = [(s.id, s.group, s.team_name, s.vong, s.diem) for s in Score.query.order_by(Score.id).all()]
    if not raw:
        return jsonify({"error": "Chưa có điểm nào để xuất."}), 404
    board = summarize(raw, {g: t["team_name"] for g, t in teams.items()})
    try:
        workbook = build_workbook(raw, board)
    except ImportError:
        return jsonify({"error": "Server thiếu thư viện openpyxl. Chạy: pip install openpyxl"}), 500
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer, as_attachment=True, download_name=f"ket_qua_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", max_age=0,
    )


@app.route("/tv")
def tv_display():
    return render_template("tv.html")


# ── Địa chỉ IP trong mạng LAN của máy chủ (để các đội biết đường vào) ──
@app.route("/network-info")
def network_info():
    try:
        result = subprocess.run(
            ["ip", "-4", "-o", "addr", "show", "scope", "global"],
            capture_output=True,
            text=True,
            timeout=3,
            check=True
        )

        ips = []
        for line in result.stdout.splitlines():
            parts = line.split()
            interface = parts[1]
            ip_cidr = parts[3]
            ip = ip_cidr.split("/")[0]

            if interface != "tailscale0" and not ip.startswith("100."):
                ips.append(ip)

        return jsonify({"lan_ips": ips})

    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ── API: Lấy dữ liệu Vòng 1 ───────────────────────────────────────
@app.route("/api/vong1/data")
def vong1_data():
    """Trả về toàn bộ bộ câu hỏi vòng 1 (file JSON thô)."""
    json_path = os.path.join(app.static_folder, "games", "vong_1.json")
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    return jsonify({
        "questions": data,
        "time": VONG1_TIME,
        "points_per_correct": POINTS_PER_CORRECT,
        "skip_penalty": SKIP_PENALTIES[1],
    })


# ── API: Lấy dữ liệu Vòng 2 ───────────────────────────────────────
@app.route("/api/vong2/data")
def vong2_data():
    """Trả về các câu ghép đôi của vòng 2 (file JSON thô), thời gian và điểm mỗi cặp."""
    return jsonify({
        "boards": _load_json("vong_2.json"),
        "time": VONG2_TIME,
        "points_per_pair": POINTS_PER_PAIR,
        "skip_penalty": SKIP_PENALTIES[2],
    })


# ── API: Lấy dữ liệu Vòng 3 ───────────────────────────────────────
@app.route("/api/vong3/data")
def vong3_data():
    """Trả về các câu xếp dòng thời gian của vòng 3 (file JSON thô), thời gian và điểm mỗi vị trí."""
    return jsonify({
        "boards": _load_json("vong_3.json"),
        "time": VONG3_TIME,
        "points_per_slot": POINTS_PER_SLOT,
        "skip_penalty": SKIP_PENALTIES[3],
    })


# ── API: Lấy dữ liệu Vòng 4 ───────────────────────────────────────
@app.route("/api/vong4/data")
def vong4_data():
    """Trả về các thẻ phân loại của vòng 4 (file JSON thô), thời gian, điểm cộng và điểm trừ."""
    data = _load_json("vong_4.json")
    data.update({
        "time": VONG4_TIME,
        "points_right": POINTS_V4_RIGHT,
        "points_wrong": POINTS_V4_WRONG,
        "skip_penalty": SKIP_PENALTIES[4],
    })
    return jsonify(data)


# ── API: Bảng điểm ─────────────────────────────────────────────────
@app.route("/api/scores")
def get_scores():
    scores = Score.query.order_by(Score.diem.desc()).all()
    return jsonify([s.to_dict() for s in scores])


# ── Socket events cho đội chơi ──────────────────────────────────────
@socketio.on("get_state")
def handle_get_state():
    """Gửi trạng thái hiện tại cho client vừa kết nối."""
    emit("state_update", _state_payload())


@socketio.on("select_group")
def handle_select_group(data):
    """Đội chơi chọn nhóm (Nhóm 1/2/4/5)."""
    group = data.get("group")

    if _game_locked():
        emit("error_message", {"message": "Trò chơi đã bắt đầu, không thể tham gia nữa!"})
        return

    if group not in ALLOWED_GROUPS:
        emit("error_message", {"message": "Nhóm không hợp lệ!"})
        return

    if any(t["sid"] == request.sid for t in teams.values()):
        emit("error_message", {"message": "Bạn đã chọn một nhóm rồi!"})
        return

    if group in teams:
        emit("error_message", {"message": f"{group} đã có đội khác chọn rồi!"})
        return

    teams[group] = {
        "team_name": group,  # Mặc định tên đội = tên nhóm
        "sid": request.sid,
        "confirmed": False,
    }

    emit("group_selected", {"group": group, "default_name": group})
    _broadcast_state()


@socketio.on("set_team_name")
def handle_set_team_name(data):
    """Đội chơi đặt tên đội và bấm Xác nhận."""
    group = data.get("group")
    team_name = (data.get("team_name") or "").strip()[:30]

    if group not in teams:
        emit("error_message", {"message": "Nhóm chưa được chọn!"})
        return

    if teams[group]["sid"] != request.sid:
        emit("error_message", {"message": "Bạn không có quyền đổi tên nhóm này!"})
        return

    if teams[group]["confirmed"]:
        emit("error_message", {"message": "Tên đội đã được xác nhận rồi!"})
        return

    teams[group]["team_name"] = team_name or group
    teams[group]["confirmed"] = True

    emit("name_confirmed", {
        "group": group,
        "team_name": teams[group]["team_name"],
        "game_status": game_state["status"],
    })
    _broadcast_state()


@socketio.on("disconnect")
def handle_disconnect():
    """Khi client ngắt kết nối, giải phóng nhóm (chỉ trước khi game bắt đầu).

    Từ lúc vòng 1 bắt đầu (kể cả giữa các vòng), mất mạng thoáng qua
    không được làm mất đội.
    """
    sid = request.sid
    btc_sids.discard(sid)
    if _game_locked():
        return

    for group in list(teams.keys()):
        if teams[group]["sid"] == sid:
            del teams[group]
            _broadcast_state()
            break


# ── Socket events cho BTC ───────────────────────────────────────────
@socketio.on("btc_get_state")
def handle_btc_get_state():
    """BTC lấy trạng thái hiện tại và vào phòng 'btc' để nhận điểm trực tiếp."""
    join_room("btc")
    btc_sids.add(request.sid)
    emit("state_update", _state_payload())
    emit("btc_scores", _btc_scores_payload())
    emit("tv_control", _tv_control_payload())


@socketio.on("tv_get_state")
def handle_tv_get_state():
    join_room("tv")
    emit("tv_update", _tv_payload())


@socketio.on("btc_tv_change")
def handle_btc_tv_change(data):
    if request.sid not in btc_sids:
        return {"ok": False}
    data = data if isinstance(data, dict) else {}
    view, raw_round = data.get("view"), data.get("round", 1)
    if view not in {"welcome", "rules", "guide", "time", "scores"}:
        emit("error_message", {"message": "Màn hình TV không hợp lệ!"})
        return {"ok": False}
    if raw_round == "total" and view == "scores":
        vong = "total"
    else:
        try:
            vong = int(raw_round)
        except (TypeError, ValueError):
            emit("error_message", {"message": "Vòng không hợp lệ!"})
            return {"ok": False}
        if vong not in ROUND_TIME:
            emit("error_message", {"message": "Vòng không hợp lệ!"})
            return {"ok": False}
    tv_control.update(view=view, round=vong)
    _broadcast_tv()
    return {"ok": True}


@socketio.on("btc_tv_reveal")
def handle_btc_tv_reveal(data):
    """BTC cho phép (hoặc thu lại) việc hiện điểm trên TV.
    data = {"round": 1..6 | "total", "groups": ["Nhóm 1", ...] | "all", "show": True/False}
    Hiện tất cả = groups "all"; hiện từng đội = danh sách đội. Ẩn thì ngược lại."""
    if request.sid not in btc_sids:
        return {"ok": False}
    data = data if isinstance(data, dict) else {}
    show = data.get("show") is not False
    raw_round, raw_groups = data.get("round"), data.get("groups", "all")

    if raw_round == "total":
        key = "total"
    elif isinstance(raw_round, int) and not isinstance(raw_round, bool) and raw_round in ROUND_TIME:
        key = raw_round
    else:
        emit("error_message", {"message": "Vòng không hợp lệ!"})
        return {"ok": False}

    if raw_groups == "all":
        groups = set(teams.keys())
    elif isinstance(raw_groups, list) and all(isinstance(g, str) for g in raw_groups):
        groups = {g for g in raw_groups if g in teams}
    else:
        emit("error_message", {"message": "Danh sách đội không hợp lệ!"})
        return {"ok": False}

    if show:
        # Chỉ được hiện khi điểm đã chốt: vòng đã kết thúc (tổng điểm: sau vòng 6)
        if key == "total":
            ready = game_state["status"] == "waiting" and game_state["round"] >= 6
        else:
            ready = 1 <= key <= game_state["round"] and not (game_state["status"] != "waiting" and key == game_state["round"])
        if not ready:
            emit("error_message", {"message": "Chỉ hiện điểm trên TV khi vòng đã kết thúc!"})
            return {"ok": False}
        tv_revealed.setdefault(key, set()).update(groups)
    else:
        current = tv_revealed.get(key, set())
        current.difference_update(groups)
        if not current:
            tv_revealed.pop(key, None)

    _broadcast_tv()
    return {"ok": True}


@socketio.on("btc_start_game")
def handle_btc_start_game(data=None):
    """BTC bấm bắt đầu vòng tiếp theo (chỉ khi mọi đội đã xác nhận tên).
    Khi DEBUG=True, BTC có thể gửi kèm {"vong": n} để bắt đầu một vòng bất kỳ (test nhanh)."""
    if game_state["status"] != "waiting":
        emit("error_message", {"message": "Một vòng đang diễn ra rồi!"})
        return

    if not teams:
        emit("error_message", {"message": "Chưa có đội nào tham gia!"})
        return

    pending = _pending_groups()
    if pending:
        emit("error_message", {
            "message": "Còn đội chưa xác nhận tên: " + ", ".join(pending),
        })
        return

    requested = data.get("vong") if isinstance(data, dict) else None
    if requested is not None:
        # Chọn vòng tuỳ ý chỉ dành cho chế độ DEBUG; server tự kiểm tra, không tin client
        if not DEBUG:
            emit("error_message", {"message": "Chỉ được chọn vòng tuỳ ý khi DEBUG=True!"})
            return
        if not isinstance(requested, int) or isinstance(requested, bool) or requested not in ROUND_TIME:
            emit("error_message", {"message": "Vòng không hợp lệ!"})
            return
        vong = requested
    else:
        vong = _next_round()
        if vong is None:
            emit("error_message", {"message": "Không còn vòng nào để bắt đầu!"})
            return

    game_state["run_id"] += 1
    run_id = game_state["run_id"]
    game_state["status"] = "countdown"
    game_state["round"] = vong
    game_state["countdown_deadline"] = time.monotonic() + COUNTDOWN_TIME
    finished_groups.clear()
    awarded_items.clear()
    skipped_items.clear()
    v4_results.clear()
    _v5_reset()

    # Bắt đầu vòng này từ đầu: xoá điểm của vòng này và các vòng sau
    # (chỉ có dữ liệu cũ khi chọn lại vòng ở chế độ DEBUG, nên vòng chạy bình thường không bị ảnh hưởng)
    Score.query.filter(Score.vong >= vong).delete()
    db.session.commit()
    revealed_rounds.difference_update([v for v in revealed_rounds if v >= vong])
    for k in [k for k in tv_revealed if k == "total" or k >= vong]:
        del tv_revealed[k]

    socketio.emit("game_countdown", {
        "round": vong,
        "seconds": COUNTDOWN_TIME,
        "teams": {g: teams[g]["team_name"] for g in teams},
    })
    _broadcast_state()
    _emit_btc_scores()

    # Server giữ đồng hồ: hết đếm ngược mới phát "game_started" cho mọi đội cùng lúc
    socketio.start_background_task(_run_round, run_id, vong)


@socketio.on("btc_reveal_scores")
def handle_btc_reveal_scores(data):
    """BTC công bố điểm của MỘT vòng cho các đội (mỗi vòng phải mở riêng)."""
    vong = (data or {}).get("vong")

    if game_state["status"] != "waiting":
        emit("error_message", {"message": "Chỉ công bố điểm khi vòng đã kết thúc!"})
        return

    if not isinstance(vong, int) or not (1 <= vong <= game_state["round"]):
        emit("error_message", {"message": "Vòng không hợp lệ!"})
        return

    revealed_rounds.add(vong)
    _broadcast_state()
    _emit_btc_scores()


@socketio.on("btc_reveal_final")
def handle_btc_reveal_final():
    if game_state["status"] != "waiting" or game_state["round"] < 6:
        emit("error_message", {"message": "Chỉ công bố tổng kết sau khi vòng 6 kết thúc!"})
        return
    v6["final_revealed"] = True
    _broadcast_state()
    _emit_btc_scores()
    socketio.emit("final_revealed", {"scores": _final_board()})


@socketio.on("btc_reset_game")
def handle_btc_reset_game():
    """BTC reset trò chơi: xoá đội, trạng thái vòng và toàn bộ điểm."""
    game_state["status"] = "waiting"
    game_state["round"] = 0
    game_state["round_deadline"] = 0.0
    game_state["countdown_deadline"] = 0.0
    game_state["run_id"] += 1  # huỷ đếm ngược / bộ đếm giờ đang chạy
    teams.clear()
    finished_groups.clear()
    revealed_rounds.clear()
    tv_revealed.clear()
    awarded_items.clear()
    skipped_items.clear()
    v4_results.clear()
    _v5_reset()
    v6.update(matches=[], questions=[], q=-1, phase="idle", final_revealed=False)

    Score.query.delete()
    db.session.commit()

    socketio.emit("game_reset", {"message": "Trò chơi đã được reset!"})
    _broadcast_state()
    _emit_btc_scores()


# ── Socket events cho đội chơi trong vòng ──────────────────────────
@socketio.on("round_finished")
def handle_round_finished(data):
    """Đội chơi báo đã chơi xong vòng (hết giờ hoặc trả lời hết câu)."""
    data = data or {}
    group = data.get("group")

    if game_state["status"] != "playing" or group not in teams:
        return
    if group in finished_groups:  # mỗi đội chỉ báo một lần
        return

    finished_groups.add(group)

    # Số liệu gửi riêng cho BTC (nhật ký). Số đúng/sai tính từ dữ liệu trên server cho đáng tin.
    def _int(key, lo, hi):
        try:
            return max(lo, min(hi, int(data.get(key, 0))))
        except (TypeError, ValueError):
            return lo

    correct, wrong = _round_stats(game_state["round"], group)
    socketio.emit("btc_round_finish", {
        "round": game_state["round"],
        "group": group,
        "team_name": teams[group]["team_name"],
        "reason": "complete" if data.get("reason") == "complete" else "timeout",
        "correct": correct,
        "wrong": wrong,
        "skipped": _int("skipped", 0, 1000),
        "total": _int("total", 0, 1000),
        "time_left": _int("time_left", 0, ROUND_TIME.get(game_state["round"], 0)),
    }, to="btc")

    if finished_groups.issuperset(teams):
        _end_round()  # cả bốn đội đã xong → kết thúc sớm, BTC có thể mở điểm
    else:
        _broadcast_state()


# ── Socket events cho Vòng 1 (cộng điểm) ───────────────────────────
def _skip_item_count(vong):
    """Số mục hợp lệ của vòng (câu ở vòng 1, câu lớn ở vòng 2–3, thẻ ở vòng 4) để kiểm tra chỉ số client gửi."""
    if vong == 4:
        return len(_load_json("vong_4.json")["the"])
    return len(_load_json(f"vong_{vong}.json"))


def _add_round_score(vong, group, delta):
    """Cộng/trừ thẳng vào điểm hiện có của đội ở vòng `vong` (có thể âm). Trả về tổng mới."""
    team_name = teams[group]["team_name"]
    score = Score.query.filter_by(group=group, vong=vong).first()
    if score:
        score.diem += delta
        score.team_name = team_name
    else:
        score = Score(group=group, team_name=team_name, vong=vong, diem=delta)
        db.session.add(score)
    db.session.commit()
    return score.diem


@socketio.on("vong1_correct")
def handle_vong1_correct(data):
    """Đội chơi trả lời đúng 1 câu trong vòng 1 → cộng điểm. Mỗi câu chỉ tính một lần cho mỗi đội.
    Client gửi {group, idx}: idx là số thứ tự gốc của câu trong vong_1.json."""
    data = data if isinstance(data, dict) else {}
    group, idx = data.get("group", ""), data.get("idx")

    # Chỉ nhận điểm khi vòng 1 đang diễn ra; điểm do server quyết định
    if game_state["status"] != "playing" or game_state["round"] != 1:
        return
    if group not in teams:
        return
    if not isinstance(idx, int) or isinstance(idx, bool):
        return
    try:
        size = _skip_item_count(1)
    except (OSError, ValueError):
        return
    key = (1, group, idx, 0)
    if not (0 <= idx < size) or key in awarded_items or (1, group, idx) in skipped_items:
        return
    awarded_items.add(key)

    _add_round_score(1, group, POINTS_PER_CORRECT)

    # Điểm trực tiếp chỉ BTC thấy; đội chơi chỉ thấy sau khi BTC công bố
    _emit_btc_scores()


# ── Socket events cho Vòng 2 và 3 (cộng điểm theo cặp / vị trí) ─────
def _award_items(vong, filename, items_key, points, data):
    """Đội chơi vừa xếp đúng một hoặc nhiều mục (cặp ở vòng 2, vị trí ở vòng 3) → cộng điểm.
    Client gửi {group, board, pairs: [chỉ số...]}; mỗi mục chỉ tính một lần, điểm do server quyết định."""
    data = data or {}
    group = data.get("group", "")
    board = data.get("board")
    items = data.get("pairs")

    if game_state["status"] != "playing" or game_state["round"] != vong:
        return
    if group not in teams:
        return
    if not isinstance(board, int) or isinstance(board, bool) or not isinstance(items, list):
        return

    try:
        boards = _load_json(filename)
    except (OSError, ValueError):
        return
    if not (0 <= board < len(boards)):
        return

    size = len(boards[board].get(items_key, []))
    new_items = {
        p for p in items
        if isinstance(p, int) and not isinstance(p, bool)
        and 0 <= p < size and (vong, group, board, p) not in awarded_items
    }
    if not new_items:
        return
    awarded_items.update((vong, group, board, p) for p in new_items)

    gain = len(new_items) * points
    team_name = teams[group]["team_name"]

    score = Score.query.filter_by(group=group, vong=vong).first()
    if score:
        score.diem += gain
        score.team_name = team_name
    else:
        db.session.add(Score(group=group, team_name=team_name, vong=vong, diem=gain))
    db.session.commit()

    _emit_btc_scores()


@socketio.on("vong2_correct")
def handle_vong2_correct(data):
    _award_items(2, "vong_2.json", "cap", POINTS_PER_PAIR, data)


@socketio.on("vong3_correct")
def handle_vong3_correct(data):
    _award_items(3, "vong_3.json", "su_kien", POINTS_PER_SLOT, data)


@socketio.on("vong4_answer")
def handle_vong4_answer(data):
    """Đội chơi trả lời một thẻ ở vòng 4. Server tự chấm theo đáp án trong vong_4.json
    và chỉ tính lần trả lời đầu tiên của mỗi thẻ: đúng cộng điểm, sai trừ điểm (tổng điểm vòng có thể âm)."""
    data = data or {}
    group = data.get("group", "")
    card = data.get("card")
    chosen = data.get("bin")

    if game_state["status"] != "playing" or game_state["round"] != 4:
        return
    if group not in teams:
        return
    if any(not isinstance(v, int) or isinstance(v, bool) for v in (card, chosen)):
        return
    if (group, card) in v4_results or (4, group, card) in skipped_items:
        return

    try:
        file = _load_json("vong_4.json")
        cards, n_bins = file["the"], len(file["ten_lop"])
    except (OSError, ValueError, KeyError):
        return
    if not (0 <= card < len(cards) and 0 <= chosen < n_bins):
        return

    right = cards[card]["lop"] == chosen
    v4_results[(group, card)] = right
    delta = POINTS_V4_RIGHT if right else -POINTS_V4_WRONG
    team_name = teams[group]["team_name"]

    # Cộng/trừ thẳng vào điểm hiện tại, kể cả khi đang âm (ví dụ -20 + 10 = -10)
    score = Score.query.filter_by(group=group, vong=4).first()
    if score:
        score.diem += delta
        score.team_name = team_name
    else:
        score = Score(group=group, team_name=team_name, vong=4, diem=delta)
        db.session.add(score)
    db.session.commit()

    _emit_btc_scores()
    return {"diem": score.diem}  # ack cho client: tổng điểm vòng 4 sau khi tính


# ── Socket events cho Vòng 5 ───────────────────────────────────────
def _v5_playing(group):
    return game_state["status"] == "playing" and game_state["round"] == 5 and group in teams


@socketio.on("v5_sync")
def handle_v5_sync(data):
    """Đội vừa vào vòng 5 (hoặc vừa nối lại mạng) xin trạng thái hiện tại; đồng thời vào phòng riêng của nhóm
    để nhận kết quả riêng (v5_result)."""
    data = data if isinstance(data, dict) else {}
    group = data.get("group", "")
    if not _v5_playing(group):
        return None
    join_room("grp:" + group)
    return _v5_snapshot(group)


@socketio.on("v5_bet")
def handle_v5_bet(data):
    """Chọn mức cược cho câu hiện tại (chỉ nhận trong pha cược)."""
    data = data if isinstance(data, dict) else {}
    group, bet = data.get("group", ""), data.get("bet")
    if not _v5_playing(group) or not isinstance(bet, int) or isinstance(bet, bool) or bet not in V5_BETS:
        return {"ok": False}
    with v5_lock:
        if v5["phase"] != "prep":
            return {"ok": False, "bet": v5["bets"].get(group, V5_DEFAULT_BET)}
        v5["bets"][group] = bet
    return {"ok": True, "bet": bet}


@socketio.on("v5_buzz")
def handle_v5_buzz(data):
    """Bấm chuông. Đội đến trước (đổi pha 'open' → 'answering' dưới khóa) được giữ chuông và nhận các lựa chọn."""
    data = data if isinstance(data, dict) else {}
    group = data.get("group", "")
    if not _v5_playing(group):
        return {"ok": False}
    with v5_lock:
        if v5["phase"] != "open" or group in v5["locked"]:
            return {"ok": False}
        v5["buzzer"] = group
        _v5_set_phase("answering", V5_ANSWER_SECONDS)
        options = list(v5["questions"][v5["q"]]["options"])
    _v5_broadcast()
    return {"ok": True, "options": options}


@socketio.on("v5_answer")
def handle_v5_answer(data):
    """Đội đang giữ chuông gửi đáp án (chỉ số trong danh sách lựa chọn đã nhận)."""
    data = data if isinstance(data, dict) else {}
    group, choice = data.get("group", ""), data.get("choice")
    if not _v5_playing(group) or not isinstance(choice, int) or isinstance(choice, bool):
        return {"ok": False}
    with v5_lock:
        q = v5["q"]
        n_options = len(v5["questions"][q]["options"]) if 0 <= q < len(v5["questions"]) else 0
    if not (0 <= choice < n_options):
        return {"ok": False}
    return {"ok": _v5_resolve_answer(group, choice)}


def _v6_get_match(group):
    return next((m for m in v6["matches"] if group in m["groups"]), None)


@socketio.on("v6_sync")
def handle_v6_sync(data):
    group = (data or {}).get("group")
    if group not in teams or game_state["status"] != "playing" or game_state["round"] != 6:
        return None
    join_room("grp:" + group)
    return _v6_snapshot(group)


@socketio.on("v6_action")
def handle_v6_action(data):
    data = data or {}
    group, action = data.get("group"), data.get("action")
    if not _v6_playing(group):
        return {"ok": False}
    with v6_lock:
        m = _v6_get_match(group)
        if not m:
            return {"ok": False}
        phase = v6["phase"]
        if phase == "bet" and group == m["opener"]:
            bet = data.get("bet")
            allowed = (20, 40, 60) if v6["q"] == V6_QUESTIONS else V6_BETS[:3]
            if not isinstance(bet, int) or isinstance(bet, bool) or bet not in allowed: return {"ok": False}
            m["bet"] = bet
        elif phase == "decide" and group != m["opener"] and action in ("follow", "raise", "fold"):
            # Luôn tính từ mức cược gốc của đội mở: bấm Tố nhiều lần hay đổi ý không làm mức cược leo thêm
            base = m.get("base_bet") or m["bet"]
            m["base_bet"] = base
            m["bet"] = base
            if action == "raise":
                ladder = (20, 40, 60, 100) if v6["q"] == V6_QUESTIONS else V6_BETS
                idx = ladder.index(base)
                if idx + 1 >= len(ladder): action = "follow"
                else: m["bet"] = ladder[idx + 1]
            m["choice"] = action
            m["call"] = "fold" if action == "fold" else None
        elif phase == "call" and group == m["opener"] and m["choice"] == "raise" and action in ("follow", "fold"):
            if action == "fold": m["call"] = "fold"
            else: m["call"] = "follow"
        elif phase == "answer" and action == "answer":
            choice = data.get("choice")
            if not isinstance(choice, int) or isinstance(choice, bool): return {"ok": False}
            question = v6["questions"][v6["q"]]
            if not 0 <= choice < len(m["variants"][group]["options"]): return {"ok": False}
            if group in m["answers"]: return {"ok": False}
            m["answers"][group] = choice
            m["speed"][group] = time.monotonic()
        else:
            return {"ok": False}
    _v6_broadcast()
    return {"ok": True}


def _v6_playing(group):
    return game_state["status"] == "playing" and game_state["round"] == 6 and group in teams


# ── Bỏ qua ở vòng 1–4: trừ điểm ────────────────────────────────────
@socketio.on("skip_item")
def handle_skip_item(data):
    """Đội bấm Bỏ qua ở vòng 1–4 → trừ SKIP_PENALTIES[vòng] điểm (chỉnh trong .env), có thể làm điểm âm.
    Mỗi mục chỉ trừ một lần. item = số thứ tự gốc trong file dữ liệu: vòng 1 = câu, vòng 2–3 = câu lớn
    (một lần bấm Bỏ qua ở một bảng), vòng 4 = thẻ. Câu/thẻ đã trả lời đúng rồi thì không tính bỏ qua."""
    data = data if isinstance(data, dict) else {}
    group, vong, item = data.get("group", ""), data.get("vong"), data.get("item")

    def is_int(v):
        return isinstance(v, int) and not isinstance(v, bool)

    if not is_int(vong) or vong not in SKIP_PENALTIES or not is_int(item):
        return {"ok": False}
    penalty = SKIP_PENALTIES[vong]
    if penalty <= 0 or game_state["status"] != "playing" or game_state["round"] != vong or group not in teams:
        return {"ok": False}
    try:
        size = _skip_item_count(vong)
    except (OSError, ValueError, KeyError):
        return {"ok": False}
    key = (vong, group, item)
    if not (0 <= item < size) or key in skipped_items:
        return {"ok": False}
    if (vong == 1 and (1, group, item, 0) in awarded_items) or (vong == 4 and (group, item) in v4_results):
        return {"ok": False}
    skipped_items.add(key)

    total = _add_round_score(vong, group, -penalty)
    _emit_btc_scores()
    return {"ok": True, "diem": total}


if __name__ == "__main__":
    with app.app_context():
        db.create_all()
        # Mỗi lần khởi động server là một ván mới: xoá toàn bộ điểm cũ trong database
        Score.query.delete()
        db.session.commit()

    socketio.run(
        app,
        host="0.0.0.0",
        port=int(os.getenv("PORT", "5000")),
        debug=True,
    )