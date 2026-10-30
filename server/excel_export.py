"""Tạo file Excel kết quả: gộp điểm thô thành bảng xếp hạng và dựng Workbook (không đụng tới Flask)."""

EXPORT_ROUNDS = range(1, 7)  # vòng 1..6


def summarize(raw, teams=None):
    """Gộp các dòng điểm thô thành bảng xếp hạng.

    raw   : [(id, group, team_name, vong, diem), ...]
    teams : {group: team_name} các đội đang tham gia; đội chưa có điểm vẫn có dòng (0 điểm).
    """
    board = {}

    def team_row(group, name):
        return board.setdefault(group, {"group": group, "team_name": name, "rounds": {v: 0 for v in EXPORT_ROUNDS}, "seen": set()})

    for _id, group, team_name, vong, diem in raw:
        team = team_row(group, team_name)
        team["team_name"] = team_name or team["team_name"]
        # App dùng .first() cho mỗi (đội, vòng) nên chỉ lấy dòng đầu tiên của mỗi cặp
        if vong in team["rounds"] and vong not in team["seen"]:
            team["rounds"][vong] = diem
            team["seen"].add(vong)
    for group, name in (teams or {}).items():
        team_row(group, name)["team_name"] = name  # tên đội hiện tại trong app là tên chuẩn

    rows = list(board.values())
    for team in rows:
        team["total"] = sum(team["rounds"].values())
    # Cùng luật xếp hạng như bảng tổng kết trong app: tổng, rồi vòng 6, 5, 4
    rows.sort(key=lambda r: (-r["total"], -r["rounds"][6], -r["rounds"][5], -r["rounds"][4], r["group"]))
    return rows


def build_workbook(raw, board):
    """Tạo Workbook (openpyxl) gồm 2 sheet: Tổng kết và Chi tiết."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.font, cell.fill, cell.alignment, cell.border = header_font, header_fill, center, border

    def autosize(ws, minimum=8, maximum=40):
        for col in ws.columns:
            width = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 3
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(minimum, min(maximum, width))

    # Sheet 1: tổng kết
    ws = wb.active
    ws.title = "Tổng kết"
    ws.append(["Hạng", "Nhóm", "Tên đội"] + [f"Vòng {v}" for v in EXPORT_ROUNDS] + ["Tổng"])
    style_header(ws)
    for rank, team in enumerate(board, start=1):
        ws.append([rank, team["group"], team["team_name"]] + [team["rounds"][v] for v in EXPORT_ROUNDS] + [team["total"]])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            if cell.column != 3:
                cell.alignment = center
        row[-1].font = Font(bold=True)
    ws.freeze_panes = "D2"
    autosize(ws)

    # Sheet 2: chi tiết từng dòng điểm trong database
    detail = wb.create_sheet("Chi tiết")
    detail.append(["ID", "Nhóm", "Tên đội", "Vòng", "Điểm"])
    style_header(detail)
    for row in raw:
        detail.append(list(row))
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            if cell.column != 3:
                cell.alignment = center
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    autosize(detail)
    return wb
